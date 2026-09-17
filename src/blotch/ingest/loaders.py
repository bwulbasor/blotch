"""Document ingestion.

The MVP extracts plain text and, on output, **regenerates a clean document from
the sanitised text** rather than editing the original binary in place. This is a
deliberate anti-leak decision (see README): regeneration cannot carry forward a
hidden text layer, annotation, form field, metadata entry, or embedded file that
in-place redaction would miss.

PDF/DOCX support is optional (``pip install 'blotch[docs]'``); ``.txt`` always
works. A *scanned* PDF has no text layer, so extraction returns nothing; when an
OCR engine is available (``pip install 'blotch[ocr]'`` plus Tesseract or
RapidOCR) the loaders fall back to OCR automatically. See :mod:`blotch.ingest.ocr`.
"""

from __future__ import annotations

import io
import os

_IMAGE = (".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp")
SUPPORTED = (".txt", ".md", ".text", ".csv", ".tsv", ".log", ".json", ".pdf",
             ".docx") + _IMAGE
_PLAIN = (".txt", ".md", ".text", ".csv", ".tsv", ".log", ".json")

# How the loaders treat OCR for scanned PDFs:
#   "auto"   - use OCR only when the text layer looks empty (a scan) and an
#              engine is available; otherwise return the text layer as-is.
#   "never"  - never OCR (today's behaviour): a scan yields empty text.
#   "always" - always OCR, ignoring any text layer (raises if no engine).
_OCR_MODES = ("auto", "never", "always")


def load_text(path: str, ocr: str = "auto") -> str:
    """Extract plain text from a supported document.

    Tabular (.csv/.tsv) and log/JSON files are read as plain text: tokenisation
    replaces values in place, so the delimiters and structure are preserved.
    """

    ext = os.path.splitext(path)[1].lower()
    if ext in _PLAIN:
        with open(path, encoding="utf-8", errors="replace") as fh:
            return fh.read()
    if ext == ".pdf" or ext in _IMAGE:
        with open(path, "rb") as fh:
            return extract_bytes(fh.read(), ext, ocr=ocr)
    if ext == ".docx":
        return _load_docx(path)
    raise ValueError(f"unsupported file type {ext!r}; supported: {', '.join(SUPPORTED)}")


def extract_bytes(data: bytes, ext: str, ocr: str = "auto") -> str:
    """Extract plain text from in-memory document ``data`` given its extension.

    Used by the daemon's upload endpoint so a PDF/DOCX can be reviewed without
    writing it to disk. ``ext`` is like ".pdf" / ".docx" / ".txt". For a scanned
    PDF, ``ocr`` controls the fallback (see ``_OCR_MODES``).
    """
    ext = ext.lower()
    if ocr not in _OCR_MODES:
        raise ValueError(f"ocr must be one of {_OCR_MODES}, got {ocr!r}")
    if ext in _PLAIN:
        return data.decode("utf-8", errors="replace")
    if ext in _IMAGE:
        from . import ocr as _ocr
        if ocr == "never":
            return ""  # an image has no text layer; OCR is the only reader
        return _ocr.ocr_image(data)  # raises a clear error if no engine
    if ext == ".pdf":
        return _extract_pdf_bytes(data, ocr)
    if ext == ".docx":
        try:
            import docx
        except Exception as exc:
            raise RuntimeError("DOCX support needs the 'docs' extra") from exc
        return "\n".join(p.text for p in docx.Document(io.BytesIO(data)).paragraphs)
    raise ValueError(f"unsupported file type {ext!r}")


def _extract_pdf_bytes(data: bytes, ocr: str) -> str:
    """Read a PDF's text layer, falling back to OCR per the ``ocr`` mode.

    In ``auto`` mode this is *per-page hybrid*: only the pages whose own text
    layer is empty get OCR'd, and their recovered text is spliced back in page
    order. So a 200-page report with three scanned inserts only rasterises those
    three pages, and a fully scanned PDF OCRs everything.
    """

    from . import ocr as _ocr

    if ocr == "always":
        return _ocr.ocr_pdf(data)

    try:
        from pypdf import PdfReader
    except Exception as exc:
        raise RuntimeError("PDF support needs the 'docs' extra") from exc
    page_texts = [(p.extract_text() or "") for p in PdfReader(io.BytesIO(data)).pages]

    if ocr == "never" or not _ocr.ocr_available():
        return "\n".join(page_texts)

    # Pages that look empty on their own are candidates for OCR.
    scanned = [i for i, t in enumerate(page_texts) if _ocr.looks_scanned(t, 1)]
    if not scanned:
        return "\n".join(page_texts)
    recovered = _ocr.ocr_pdf_pages(data, scanned)
    for i, text in recovered.items():
        page_texts[i] = text
    return "\n".join(page_texts)


def _load_docx(path: str) -> str:  # pragma: no cover - optional dependency
    try:
        import docx
    except Exception as exc:
        raise RuntimeError(
            "DOCX support needs the optional 'docs' extra: pip install 'blotch[docs]'"
        ) from exc
    document = docx.Document(path)
    return "\n".join(p.text for p in document.paragraphs)
